from io import BytesIO
from uuid import UUID

import pandas as pd
from docx import Document
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.enums import ReportStatus
from app.models.file_asset import FileAsset
from app.models.report import Report
from app.services.aggregation_service import apply_date_filter, compute_engagements, top_ranked_posts
from app.services.audit_service import create_audit_log
from app.services.r2_service import R2Service


async def generate_report_background(
    db: AsyncSession,
    r2_service: R2Service,
    *,
    report_id: UUID,
    file_id: UUID,
    user_id: str,
    start_date,
    end_date,
) -> None:
    report = await db.get(Report, report_id)
    file_asset = await db.get(FileAsset, file_id)
    if not report or not file_asset:
        return

    try:
        raw = await r2_service.get_object(file_asset.r2_key)
        df = pd.read_excel(BytesIO(raw)) if file_asset.original_filename.lower().endswith(('xls', 'xlsx')) else pd.read_csv(BytesIO(raw))
        filtered = apply_date_filter(df, start_date, end_date)
        ranked = top_ranked_posts(compute_engagements(filtered))

        excel_buffer = BytesIO()
        ranked.to_excel(excel_buffer, index=False)
        excel_bytes = excel_buffer.getvalue()

        doc = Document()
        doc.add_heading('Weekly Social Media Performance Report', level=1)
        for _, row in ranked.iterrows():
            doc.add_paragraph(f"Rank {row['Rank']}: {row['Post text']}")
        word_buffer = BytesIO()
        doc.save(word_buffer)

        report_key = f'reports/{report.id}/report.xlsx'
        report_doc_key = f'reports/{report.id}/report.docx'
        await r2_service.put_object(report_key, excel_bytes, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        await r2_service.put_object(report_doc_key, word_buffer.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')

        report.status = ReportStatus.COMPLETED.value
        report.output_r2_key = report_key
        await db.commit()

        await create_audit_log(db, action='REPORT_GENERATE', user_id=user_id, resource_id=str(report.id), payload={'file_id': str(file_id)})
    except Exception as exc:  # intentionally broad for task status update
        report.status = ReportStatus.FAILED.value
        await db.commit()
        await create_audit_log(db, action='REPORT_FAILED', user_id=user_id, resource_id=str(report.id), payload={'error': str(exc)})


async def create_report_request(db: AsyncSession, *, file_id: UUID, user_id: str) -> Report:
    report = Report(file_id=file_id, requested_by=user_id)
    db.add(report)
    await db.commit()
    await db.refresh(report)
    return report


async def get_report(db: AsyncSession, report_id: UUID) -> Report | None:
    return await db.get(Report, report_id)


async def summary_stats(db: AsyncSession) -> dict[str, int]:
    uploads = await db.scalar(select(func.count(FileAsset.id)))
    reports = await db.scalar(select(func.count(Report.id)))
    completed = await db.scalar(select(func.count(Report.id)).where(Report.status == ReportStatus.COMPLETED.value))
    failed = await db.scalar(select(func.count(Report.id)).where(Report.status == ReportStatus.FAILED.value))
    return {
        'total_uploads': uploads or 0,
        'total_reports': reports or 0,
        'completed_reports': completed or 0,
        'failed_reports': failed or 0,
    }
